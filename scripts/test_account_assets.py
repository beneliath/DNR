#!/usr/bin/env python3
"""Generate original, fictional demo downloads; no production assets are used."""
import argparse
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer


def build(destination):
    destination.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    for name, title, sections in [
        ('speaker-notes', 'Hope, History, and Meaningful Conversations', [
            ('Purpose', 'Practice listening carefully and connecting biblical study with everyday community life.'),
            ('Session outline', 'Welcome and context; explore a passage together; discuss how history shapes our questions; identify practical next steps.'),
            ('Discussion questions', 'What surprised you? What would you ask next? How can your community make space for a respectful conversation?'),
            ('Host checklist', 'Confirm microphone and seating. Offer printed resources. Allow fifteen minutes for questions and a host debrief.'),
        ]),
        ('host-guide', 'Hosting a Community Teaching Weekend', [
            ('Before the event', 'Confirm the venue, attendance estimate, presentation times, local transport, lodging, and resource table.'),
            ('On the day', 'Meet the speaker, test projection and audio, share QR resources, and record actual attendance.'),
            ('After the event', 'Record the host conversation in Chron, complete follow-up tasks, and enter actual financial receipts.'),
        ]),
    ]:
        story = [Paragraph('TEST ACCOUNT - FICTIONAL TRAINING RESOURCE', styles['Heading3']),
                 Paragraph(title, styles['Title']), Spacer(1, 18)]
        for heading, body in sections:
            story += [Paragraph(heading, styles['Heading2']), Paragraph(body, styles['BodyText']), Spacer(1, 14)]
        story += [Spacer(1, 18), Paragraph('You may download, replace, or delete these resources while exploring the app.', styles['BodyText'])]
        SimpleDocTemplate(str(destination / (name + '.pdf')), pagesize=letter,
                          rightMargin=54, leftMargin=54, topMargin=48, bottomMargin=48).build(story)
    deck = Presentation()
    for title, lines in [
        ('Hope, History, and Community', ['Test Account', 'Fictional presentation for testing downloads and QR resources']),
        ('Start with listening', ['Ask an open question', 'Make room for different experiences', 'Record a useful next step']),
        ('Study and discuss', ['Read a passage in context', 'Connect the historical setting', 'Explore questions with care']),
        ('Plan a teaching weekend', ['Confirm host and on-site contacts', 'Review the schedule and travel', 'Prepare notes and resource links']),
        ('After the event', ['Record actual attendance', 'Complete tasks and financial closeout', 'Follow up with the host']),
        ('Try it yourself', ['Edit the presentation', 'Download or replace this slidedeck', 'Explore the speaker links and QR codes']),
    ]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = '\n'.join(lines)
        for paragraph in slide.placeholders[1].text_frame.paragraphs:
            paragraph.font.size = Pt(25)
    deck.save(destination / 'community-teaching.pptx')
    for index, initials in enumerate(['EL', 'NB', 'MC', 'JR', 'AL', 'DP']):
        im = Image.new('RGB', (320, 320), ['#1e4f66', '#6d426d', '#37644f'][index % 3])
        draw = ImageDraw.Draw(im)
        draw.ellipse((85, 45, 235, 195), fill='#e2eced')
        draw.rounded_rectangle((40, 180, 280, 320), 80, fill='#e2eced')
        draw.text((145, 115), initials, fill='#173d4a')
        draw.text((110, 280), 'DEMO AVATAR', fill='#173d4a')
        im.save(destination / f'avatar-{index}.png')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('destination', type=Path)
    build(parser.parse_args().destination)
